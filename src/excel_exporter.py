import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def export_comps_to_excel(valuations: list, filename: str = "comps_output.xlsx"):
    """Exports structured valuation summaries with CAGR and summary statistics rows."""
    print("Generating professional Excel comps deliverable with summary metrics...")
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Trading Comps"
    
    ws.views.sheetView[0].showGridLines = True
    
    # Styling definitions
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    summary_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=11)
    bold_data_font = Font(name="Calibri", size=11, bold=True)
    
    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
    )
    summary_border = Border(
        left=Side(style='thin', color='000000'), right=Side(style='thin', color='000000'),
        top=Side(style='thin', color='000000'), bottom=Side(style='double', color='000000')
    )

    # Title block
    ws['B2'] = "Automated Trading Comparables Analysis"
    ws['B2'].font = Font(name="Calibri", size=16, bold=True, color="1F4E78")
    ws['B3'] = "Institutional-Grade Multi-Agent Valuation Model with Growth Metrics"
    ws['B3'].font = Font(name="Calibri", size=11, italic=True, color="595959")

    # Table Headers
    headers = [
        "Ticker", "Company Name", "Share Price ($)", "Market Cap ($M)", 
        "Enterprise Value ($M)", "LTM Revenue ($M)", "Rev CAGR", "EBITDA ($M)", 
        "EV / Rev", "EV / EBITDA", "P/E Ratio", "Debt / EBITDA"
    ]
    
    header_row = 5
    for col_idx, header in enumerate(headers, start=2):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    # Data Rows
    start_row = 6
    current_row = start_row
    for val in valuations:
        # Calculate 2-year Revenue CAGR if historical data exists
        cagr = 0.0
        if val.revenue_two_years_ago > 0 and val.revenue > 0:
            cagr = (val.revenue / val.revenue_two_years_ago) ** (1 / 2) - 1

        row_data = [
            val.ticker,
            val.company_name,
            val.share_price,
            val.market_cap,
            val.enterprise_value,
            val.revenue,
            cagr,
            val.ebitda,
            val.ev_to_revenue,
            val.ev_to_ebitda,
            val.pe_ratio,
            val.debt_to_ebitda
        ]
        
        for col_idx, value in enumerate(row_data, start=2):
            cell = ws.cell(row=current_row, column=col_idx, value=value)
            cell.font = data_font
            cell.border = thin_border
            
            if col_idx in [2, 3]:
                cell.alignment = align_left
            else:
                cell.alignment = align_right
                
            # Formats
            if col_idx in [4, 5, 6, 7, 9]:
                cell.number_format = "$#,##0.0"
            elif col_idx == 8:
                cell.number_format = "0.0%"
            elif col_idx in [10, 11, 12, 13]:
                cell.number_format = "0.0x" if col_idx != 12 else "0.0"
                
        current_row += 1

    end_data_row = current_row - 1

    # Summary Statistics Rows (Mean, Median, Min, Max)
    summary_metrics = [
        ("Mean", "=AVERAGE(D{0}:D{1})", "AVERAGE"),
        ("Median", "=MEDIAN(D{0}:D{1})", "MEDIAN"),
        ("Min", "=MIN(D{0}:D{1})", "MIN"),
        ("Max", "=MAX(D{0}:D{1})", "MAX")
    ]

    for label, formula_template, func_name in summary_metrics:
        ws.cell(row=current_row, column=2, value=label).font = bold_data_font
        ws.cell(row=current_row, column=2).fill = summary_fill
        ws.cell(row=current_row, column=2).alignment = align_left
        ws.cell(row=current_row, column=3, value="")
        ws.cell(row=current_row, column=3).fill = summary_fill

        # Apply formulas across numeric columns (Cols 4 through 13)
        for col_idx in range(4, 14):
            col_letter = get_column_letter(col_idx)
            cell = ws.cell(row=current_row, column=col_idx)
            cell.value = f"={func_name}({col_letter}{start_row}:{col_letter}{end_data_row})"
            cell.font = bold_data_font
            cell.fill = summary_fill
            cell.alignment = align_right
            cell.border = summary_border
            
            # Formats for summary row
            if col_idx in [4, 5, 6, 7, 9]:
                cell.number_format = "$#,##0.0"
            elif col_idx == 8:
                cell.number_format = "0.0%"
            elif col_idx in [10, 11, 12, 13]:
                cell.number_format = "0.0x" if col_idx != 12 else "0.0"

        current_row += 1

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(filename)
    print(f"Successfully exported comps with CAGR and summary stats to {filename}!")